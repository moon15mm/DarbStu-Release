# -*- coding: utf-8 -*-
"""
attendance_report.py — تقرير الرصد: غيابٌ وتأخرٌ ودقائقه بين تاريخين.

لماذا وحدة مستقلة: تقارير `report_builder` كلها **غيابٌ وحده**، ووكيل
شؤون الطلاب يرصد الاثنين معاً — من غاب خمساً ومن تأخر عشراً ومن أهدر
ساعتين من الحصة الأولى. وجمعُهما في استعلامٍ واحد يُظهر ما لا يُظهره
التقريران منفصلين: الطالب الذي لا يبلغ عتبة الغياب ولا عتبة التأخر لكنه
يبلغ مجموعهما.

وحدةٌ جديدة لا إضافةٌ إلى `report_builder`: ذاك بلغ ٨٠٠ سطر، والجديد
يحتاج اختباراً مستقلاً — انظر `tools/test_attendance_report.py`.

مصادر البيانات (كلاهما في `absences.db`):
  • `absences`  — فريدٌ بـ(تاريخ، فصل، طالب)  ⇒ العدّ «أيام غياب»
  • `tardiness` — فريدٌ بـ(تاريخ، طالب) و`minutes_late` ⇒ «أيام تأخر»
  • `excuses`   — للاستبعاد الاختياري
  • `exempted_students` — يُستبعَد دائماً (مستثنى من الرصد أصلاً)
"""
import calendar
import datetime
import sqlite3

from constants import DB_PATH

# ── الأسبوع الدراسي السعودي: الأحد → الخميس ──────────────────────────
# ‏Python يبدأ الأسبوع بالاثنين (weekday(): الاثنين=0 … الأحد=6)، فحساب
# «هذا الأسبوع» بالافتراضي يضع الأحد في الأسبوع الماضي — ويوم الأحد
# أثقل أيام الغياب، فتسقط ذروةُ الأسبوع من تقريره.
_SUNDAY = 6

GROUPS = ("student", "class", "day")
SORTS = ("absence", "tardy", "minutes", "total", "name")


def _con():
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    return con


# ══════════════════════════════════════════════════════════════════
#  المدى الزمني
# ══════════════════════════════════════════════════════════════════
def resolve_range(preset="month", start=None, end=None, today=None):
    """
    يحوّل اختصاراً زمنياً إلى (من، إلى) بصيغة YYYY-MM-DD.

    المنطق هنا لا في الواجهة: الواجهة تُحسب بتوقيت جهاز المستعرِض، وقد
    يكون مضبوطاً خطأً أو في منطقة أخرى، فيختلف «اليوم» عن يوم المدرسة.
    """
    t = today or datetime.date.today()
    if isinstance(t, str):
        t = datetime.date.fromisoformat(t)

    if preset == "custom":
        s = start or t.isoformat()
        e = end or t.isoformat()
        if s > e:                       # مقلوبان ⇒ نصحّح بدل أن نُرجع فراغاً
            s, e = e, s
        return s, e
    if preset == "today":
        return t.isoformat(), t.isoformat()
    if preset == "yesterday":
        y = t - datetime.timedelta(days=1)
        return y.isoformat(), y.isoformat()
    if preset == "week":
        back = (t.weekday() - _SUNDAY) % 7
        s = t - datetime.timedelta(days=back)
        return s.isoformat(), t.isoformat()
    if preset == "last_week":
        back = (t.weekday() - _SUNDAY) % 7
        this_sun = t - datetime.timedelta(days=back)
        s = this_sun - datetime.timedelta(days=7)
        return s.isoformat(), (this_sun - datetime.timedelta(days=1)).isoformat()
    if preset == "last_month":
        first = t.replace(day=1)
        prev_end = first - datetime.timedelta(days=1)
        return prev_end.replace(day=1).isoformat(), prev_end.isoformat()
    if preset == "term":
        # الفصل الدراسي تقريباً: من أول أغسطس إن كنا في النصف الثاني من
        # السنة، وإلا من أول يناير. تقريبٌ مقصود — لا تقويم رسمي في
        # النظام، والوكيل يضبط المدى يدوياً إن أراد الدقة.
        s = t.replace(month=8, day=1) if t.month >= 8 else t.replace(month=1, day=1)
        return s.isoformat(), t.isoformat()
    # الافتراضي: هذا الشهر
    return t.replace(day=1).isoformat(), t.isoformat()


def _days_label(start, end):
    s = datetime.date.fromisoformat(start)
    e = datetime.date.fromisoformat(end)
    return (e - s).days + 1


# ══════════════════════════════════════════════════════════════════
#  التقرير
# ══════════════════════════════════════════════════════════════════
def _norm_cid(cid):
    """معرّف الفصل مُطبَّعاً — نفس ما تفعله كل تقارير `report_builder`."""
    try:
        from database import normalize_legacy_class_id
        return str(normalize_legacy_class_id(cid) or "").strip()
    except Exception:
        return str(cid or "").strip()


def class_level(cid):
    """
    الصفّ من معرّف الفصل: `1-101` ⇒ `1`. وهو ما يُجمَع عليه «كل الشعب».

    معرّفات نور صفٌّ ثم شرطة ثم الشعبة، فالجزء الأول هو الصفّ. وما لا
    يتبع الصيغة يُرجَع كما هو فلا يسقط من التقارير.
    """
    c = _norm_cid(cid)
    return c.split("-", 1)[0] if "-" in c else c


def _cls_match(cid, class_id=None, level=None):
    """هل يدخل هذا الفصل في التصفية؟ فصلٌ بعينه، أو صفٌّ بكل شعبه."""
    c = _norm_cid(cid)
    if class_id:
        return c == _norm_cid(class_id)
    if level:
        return class_level(c) == str(level).strip()
    return True


def build_report(start, end, include_absence=True, include_tardy=True,
                 class_id=None, group_by="student", min_absence=0,
                 min_tardy=0, min_minutes=0, exclude_excused=False,
                 sort="absence", limit=0, roster=None, level=None):
    """
    يبني التقرير المجمَّع. يُرجع قاموساً جاهزاً للعرض أو الطباعة.

    `class_id` فصلٌ بعينه، و`level` صفٌّ كامل بكل شعبه (مثلاً كل شعب
    «ثاني ثانوي»). والفصل أخصّ فيُقدَّم عند وجودهما معاً.

    `min_*` عتباتٌ **تُطبَّق بعد الجمع**: «أرني من غاب ٥ أيام فأكثر» هو
    سؤال الوكيل الأول، وتصفيتُه في الواجهة تعني نقل آلاف الصفوف ليُرمى
    أكثرها.
    """
    if group_by not in GROUPS:
        group_by = "student"
    if sort not in SORTS:
        sort = "absence"
    if not include_absence and not include_tardy:
        include_absence = include_tardy = True      # لا تقرير بلا مصدر

    # ⚠️ سجلات البصمة تُكتب بـ`class_id` **فارغ** (اسم الفصل وحده) —
    # انظر `poller.process_pending`. فترشيحٌ بالفصل أو الصفّ لا يُطابق
    # منها شيئاً، وتظهر «ثاني ثانوي» بطالبةٍ واحدة سُجّلت يدوياً وحدها.
    # فنعود إلى كشف الطلاب: الطالب يعرف فصله وإن لم يعرفه السجلّ.
    # (وهذا يُنقذ السجلات القائمة كلها، لا القادمة فقط.)
    _by_stu = (roster or {}).get("by_student") or {}

    def _row_cid(row):
        c = _norm_cid(row.get("class_id"))
        if c:
            return c
        return _norm_cid(_by_stu.get(str(row.get("student_id") or ""), ""))

    def _row_match(row):
        return _cls_match(_row_cid(row), class_id, level)

    con = _con()
    try:
        cur = con.cursor()

        # المستثنون: خارج الرصد أصلاً، فلا يُحسبون ولا يُعرضون
        exempt = {str(r[0]) for r in
                  cur.execute("SELECT student_id FROM exempted_students")}

        # ⚠️ لا تُرشَّح الفصول في SQL. المخزَّن في الجداول قد يكون معرّفاً
        # قديماً من نور (`1314-101`) بينما القائمة تعرض المُطبَّع
        # (`1-101`) — فالمقارنة الحرفية لا تُطابق شيئاً، ويظهر التقرير
        # فارغاً مع أن المتأخرات موجودات. (كل تقارير `report_builder`
        # تُطبّع بعد الجلب، وهذا ما فاتني هنا.)
        # فنجلب ثم نُطبّع ثم نُرشّح في بايثون — والتطبيع نفسه يوحّد
        # المعرّفين القديم والجديد فلا ينشطر الفصل الواحد مجموعتين.
        where_cls = ""
        p_cls = []

        # ── الغياب ────────────────────────────────────────────
        absences = []
        if include_absence:
            q = ("SELECT date, student_id, student_name, class_id, class_name "
                 "FROM absences WHERE date BETWEEN ? AND ?" + where_cls)
            # يُحوَّل إلى قاموس **قبل** الترشيح: `sqlite3.Row` بلا `.get`
            absences = [d for d in (dict(r) for r in
                                    cur.execute(q, [start, end] + p_cls))
                        if _row_match(d)]

        # ── الأعذار (للاستبعاد) ───────────────────────────────
        excused = set()
        if exclude_excused:
            q = "SELECT date, student_id FROM excuses WHERE date BETWEEN ? AND ?"
            excused = {(r["date"], str(r["student_id"]))
                       for r in cur.execute(q, [start, end])}

        # ── التأخر ────────────────────────────────────────────
        tardies = []
        if include_tardy:
            q = ("SELECT date, student_id, student_name, class_id, class_name,"
                 " COALESCE(minutes_late,0) AS mins "
                 "FROM tardiness WHERE date BETWEEN ? AND ?" + where_cls)
            tardies = [d for d in (dict(r) for r in
                                   cur.execute(q, [start, end] + p_cls))
                       if _row_match(d)]

        # أيام الدراسة الفعلية: التواريخ التي سُجّل فيها شيء. أدقّ من
        # عدّ أيام التقويم — العطل والإجازات لا سجلّ فيها.
        school_days = len({a["date"] for a in absences} |
                          {t["date"] for t in tardies})
    finally:
        con.close()

    # ── التجميع ───────────────────────────────────────────────
    agg = {}

    def _slot(key, name, cls):
        if key not in agg:
            agg[key] = {"key": key, "name": name, "class_name": cls,
                        "absence": 0, "tardy": 0, "minutes": 0,
                        "excused": 0, "dates_abs": [], "dates_tardy": []}
        return agg[key]

    def _keyof(row):
        if group_by == "class":
            # المُطبَّع لا الخام: القديم والجديد لفصلٍ واحد يجب أن
            # يجتمعا في صفٍّ واحد لا صفَّين
            return (_row_cid(row) or "—",
                    row.get("class_name") or "—", "")
        if group_by == "day":
            return (row["date"], row["date"], "")
        sid = str(row.get("student_id") or "")
        return (sid, row.get("student_name") or sid,
                row.get("class_name") or "")

    for a in absences:
        sid = str(a.get("student_id") or "")
        if sid in exempt:
            continue
        k, nm, cls = _keyof(a)
        s = _slot(k, nm, cls)
        if exclude_excused and (a["date"], sid) in excused:
            s["excused"] += 1
            continue
        s["absence"] += 1
        s["dates_abs"].append(a["date"])

    for t in tardies:
        sid = str(t.get("student_id") or "")
        if sid in exempt:
            continue
        k, nm, cls = _keyof(t)
        s = _slot(k, nm, cls)
        s["tardy"] += 1
        s["minutes"] += int(t.get("mins") or 0)
        s["dates_tardy"].append(t["date"])

    rows = list(agg.values())

    # ── العتبات ───────────────────────────────────────────────
    # تُطبَّق معاً لا بالتناوب: من طلب «غياب ≥٥ وتأخر ≥٣» يريد من جمعهما
    if min_absence:
        rows = [r for r in rows if r["absence"] >= min_absence]
    if min_tardy:
        rows = [r for r in rows if r["tardy"] >= min_tardy]
    if min_minutes:
        rows = [r for r in rows if r["minutes"] >= min_minutes]

    # ── النسبة: المقام يختلف باختلاف التجميع ──────────────────
    # القسمة على أيام الدراسة وحدها صحيحةٌ للطالب فقط. وتطبيقها على
    # الفصل أعطى **٢٥٣٪** في تجربة الديمو — ورقمٌ كهذا يُفقد الثقة
    # بالتقرير كله لا بعموده وحده.
    #   طالب → أيام الدراسة          (كم يوماً من أيامه غابه)
    #   فصل  → أيام الدراسة × عدده   (كم حصةَ حضورٍ ضاعت من الفصل)
    #   يوم  → عدد طلاب المدرسة      (كم نسبة الغائبين ذلك اليوم)
    # وبلا كشفٍ للطلاب لا مقام للفصل واليوم، فتُحذف النسبة ولا تُلفَّق.
    roster = roster or {}
    r_total = int(roster.get("total") or 0)
    r_by_class = roster.get("by_class") or {}

    for r in rows:
        r["total"] = r["absence"] + r["tardy"]
        if group_by == "student":
            denom = school_days
        elif group_by == "class":
            denom = school_days * int(r_by_class.get(r["key"], 0) or 0)
        else:
            denom = r_total
        r["pct"] = round(r["absence"] * 100.0 / denom, 1) if denom > 0 else None
        r["last_absence"] = max(r["dates_abs"]) if r["dates_abs"] else ""
        r["last_tardy"] = max(r["dates_tardy"]) if r["dates_tardy"] else ""
        r["avg_minutes"] = round(r["minutes"] / r["tardy"], 1) if r["tardy"] else 0.0
        r.pop("dates_abs", None)
        r.pop("dates_tardy", None)

    if sort == "name":
        rows.sort(key=lambda r: r["name"])
    elif group_by == "day":
        rows.sort(key=lambda r: r["key"])        # الأيام بترتيبها دائماً
    else:
        rows.sort(key=lambda r: (-r.get(sort, 0), r["name"]))

    if limit and limit > 0:
        rows = rows[:limit]

    return {
        "ok": True,
        "from": start, "to": end,
        "span_days": _days_label(start, end),
        "school_days": school_days,
        "group_by": group_by, "sort": sort,
        "include_absence": include_absence, "include_tardy": include_tardy,
        "exclude_excused": exclude_excused,
        # نسبةٌ بلا مقام لا تُعرض: عمودٌ من الشُّرَط أسوأ من غيابه
        "has_pct": any(r.get("pct") is not None for r in rows),
        "filters": {"class_id": class_id or "", "level": level or "",
                    "min_absence": min_absence,
                    "min_tardy": min_tardy, "min_minutes": min_minutes},
        "totals": {
            "rows": len(rows),
            "absence": sum(r["absence"] for r in rows),
            "tardy": sum(r["tardy"] for r in rows),
            "minutes": sum(r["minutes"] for r in rows),
            "excused": sum(r["excused"] for r in rows),
        },
        "rows": rows,
    }


# ══════════════════════════════════════════════════════════════════
#  المنضبطون
# ══════════════════════════════════════════════════════════════════
def build_good_standing(start, end, students, check_absence=True,
                        check_tardy=True, max_absence=0, max_tardy=0,
                        excused_ok=True, class_id=None, sort="name",
                        limit=0, level=None):
    """
    الطلاب الذين لم يغيبوا ولم يتأخروا في المدى (أو ضمن هامشٍ مسموح).

    ⚠️ **لا يُستخرج هذا بقلب تقرير الرصد.** الرصد يقرأ جدولَي الغياب
    والتأخر، ومن لم يغب قط **لا سجلّ له فيهما إطلاقاً** — فقلبُ قائمته
    يُسقط أنضباط الطلاب جميعاً. المصدر هنا هو **الكشف**، والجدولان
    يُستعملان للاستبعاد لا للإدراج.

    `excused_ok`: الغياب بعذرٍ لا يكسر الانضباط. مفعَّلٌ افتراضياً —
    من مرض بإذن المدرسة لا يُعاقَب بحرمانه من قائمة التكريم.

    `max_*`: هامشٌ يختاره الوكيل. صفرٌ = الانضباط التام، وهو الافتراض؛
    لكن «لا غياب إطلاقاً» قد يُفرّغ القائمة في مدرسة كبيرة فيُسمح برفعه.
    """
    if not check_absence and not check_tardy:
        check_absence = check_tardy = True

    con = _con()
    try:
        cur = con.cursor()
        exempt = {str(r[0]) for r in
                  cur.execute("SELECT student_id FROM exempted_students")}

        excused = set()
        if excused_ok:
            excused = {(r["date"], str(r["student_id"])) for r in cur.execute(
                "SELECT date, student_id FROM excuses WHERE date BETWEEN ? AND ?",
                [start, end])}

        abs_rows = [dict(r) for r in cur.execute(
            "SELECT date, student_id FROM absences WHERE date BETWEEN ? AND ?",
            [start, end])]
        tdy_rows = [dict(r) for r in cur.execute(
            "SELECT student_id, COALESCE(minutes_late,0) AS mins "
            "FROM tardiness WHERE date BETWEEN ? AND ?", [start, end])]

        school_days = len({a["date"] for a in abs_rows}) or len(
            {r["date"] for r in cur.execute(
                "SELECT DISTINCT date FROM tardiness WHERE date BETWEEN ? AND ?",
                [start, end])})
    finally:
        con.close()

    n_abs, n_exc, n_tdy, n_min = {}, {}, {}, {}
    for a in abs_rows:
        sid = str(a["student_id"])
        if excused_ok and (a["date"], sid) in excused:
            n_exc[sid] = n_exc.get(sid, 0) + 1
        else:
            n_abs[sid] = n_abs.get(sid, 0) + 1
    for t in tdy_rows:
        sid = str(t["student_id"])
        n_tdy[sid] = n_tdy.get(sid, 0) + 1
        n_min[sid] = n_min.get(sid, 0) + int(t.get("mins") or 0)

    rows, scanned = [], 0
    for s in students or []:
        sid = str(s.get("id") or "")
        if not sid or sid in exempt:
            continue
        # يُطبَّع هنا أيضاً: كشف الطلاب قد يحمل معرّفاً، والمختار آخر
        if not _cls_match(s.get("class_id"), class_id, level):
            continue
        scanned += 1
        a, t = n_abs.get(sid, 0), n_tdy.get(sid, 0)
        if check_absence and a > max_absence:
            continue
        if check_tardy and t > max_tardy:
            continue
        rows.append({
            "key": sid, "name": s.get("name") or sid,
            "class_name": s.get("class_name") or "",
            "absence": a, "tardy": t,
            "minutes": n_min.get(sid, 0), "excused": n_exc.get(sid, 0),
            "total": a + t, "pct": None,
            "perfect": (a == 0 and t == 0),
        })

    if sort == "class":
        rows.sort(key=lambda r: (r["class_name"], r["name"]))
    elif sort == "clean":
        rows.sort(key=lambda r: (r["total"], r["minutes"], r["name"]))
    else:
        rows.sort(key=lambda r: r["name"])
    if limit and limit > 0:
        rows = rows[:limit]

    return {
        "ok": True, "kind": "good",
        "from": start, "to": end,
        "span_days": _days_label(start, end),
        "school_days": school_days,
        "group_by": "student", "sort": sort,
        "include_absence": check_absence, "include_tardy": check_tardy,
        "exclude_excused": False, "has_pct": False,
        "excused_ok": excused_ok,
        "filters": {"class_id": class_id or "", "level": level or "",
                    "max_absence": max_absence,
                    "max_tardy": max_tardy},
        "totals": {
            "rows": len(rows), "scanned": scanned,
            "perfect": sum(1 for r in rows if r["perfect"]),
            "absence": sum(r["absence"] for r in rows),
            "tardy": sum(r["tardy"] for r in rows),
            "minutes": sum(r["minutes"] for r in rows),
            "excused": sum(r["excused"] for r in rows),
        },
        "rows": rows,
    }


# ══════════════════════════════════════════════════════════════════
#  العرض: الطباعة والتصدير
# ══════════════════════════════════════════════════════════════════
_GROUP_AR = {"student": "الطالب", "class": "الفصل", "day": "اليوم"}


def _esc(s):
    return (str(s if s is not None else "")
            .replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def _fmt_minutes(m):
    """٩٥ دقيقة ⇒ «١ س ٣٥ د» — الوكيل يقرأ الساعات لا الدقائق الخام."""
    m = int(m or 0)
    if m < 60:
        return "%d د" % m
    return "%d س %d د" % (m // 60, m % 60)


def columns(rep):
    """أعمدة التقرير حسب ما اختير — ومصدرٌ واحد للطباعة والتصدير معاً."""
    if rep.get("kind") == "good":
        # قائمة تكريم: تُقرأ بأسماء الطلاب لا بأرقام غيابهم، فالأعمدة
        # الرقمية تظهر **فقط** حين يُسمح بهامش — وإلا كانت أصفاراً تملأ
        # الورقة بلا معنى
        cols = [("#", "idx"), ("الطالب", "name"), ("الفصل", "class_name")]
        f = rep.get("filters", {})
        if rep["include_absence"] and f.get("max_absence"):
            cols.append(("أيام الغياب", "absence"))
        if rep["include_tardy"] and f.get("max_tardy"):
            cols += [("أيام التأخر", "tardy"), ("مجموع الدقائق", "minutes")]
        if rep.get("excused_ok"):
            cols.append(("غياب بعذر", "excused"))
        return cols

    cols = [("#", "idx"), (_GROUP_AR.get(rep["group_by"], "البند"), "name")]
    if rep["group_by"] == "student":
        cols.append(("الفصل", "class_name"))
    if rep["include_absence"]:
        cols.append(("أيام الغياب", "absence"))
        if rep.get("has_pct", True):
            cols.append(("النسبة %", "pct"))
        if rep["exclude_excused"]:
            cols.append(("بعذر", "excused"))
        if rep["group_by"] != "day":
            cols.append(("آخر غياب", "last_absence"))
    if rep["include_tardy"]:
        cols += [("أيام التأخر", "tardy"), ("مجموع الدقائق", "minutes"),
                 ("متوسط الدقائق", "avg_minutes")]
        if rep["group_by"] != "day":
            cols.append(("آخر تأخر", "last_tardy"))
    if rep["include_absence"] and rep["include_tardy"]:
        cols.append(("المجموع", "total"))
    return cols


def _cell(row, key, idx):
    if key == "idx":
        return str(idx)
    v = row.get(key, "")
    if key == "minutes":
        return _fmt_minutes(v)
    if key == "pct":
        # صفٌّ بلا مقام وسط صفوفٍ لها مقام (فصلٌ ليس في الكشف مثلاً):
        # «0.0%» كذبٌ صريح، والشَّرطة تقول «لا أعرف»
        return "—" if v is None else "%.1f%%" % v
    return _esc(v)


def build_print_html(rep, school_name="", title=""):
    """صفحة A4 جاهزة للطباعة. لا خطوط من الإنترنت: مدارس كثيرة بلا
    اتصالٍ سريع، واستيراد خطٍّ بعيد يؤخّر الطباعة بلا فائدة."""
    cols = columns(rep)
    head = "".join("<th>%s</th>" % _esc(h) for h, _ in cols)

    body = []
    for i, r in enumerate(rep["rows"], 1):
        tds = "".join("<td>%s</td>" % _cell(r, k, i) for _, k in cols)
        body.append("<tr>%s</tr>" % tds)
    if not body:
        body.append('<tr><td colspan="%d" class="empty">لا توجد سجلات '
                    'مطابقة لهذه الخيارات</td></tr>' % len(cols))

    f = rep["filters"]
    conds = []
    if f.get("min_absence"):
        conds.append("غياب ≥ %d" % f["min_absence"])
    if f.get("min_tardy"):
        conds.append("تأخر ≥ %d" % f["min_tardy"])
    if f.get("min_minutes"):
        conds.append("دقائق ≥ %d" % f["min_minutes"])
    if rep["exclude_excused"]:
        conds.append("بلا الغياب بعذر")
    if rep.get("kind") == "good":
        conds = []
        if rep["include_absence"]:
            conds.append("غياب لا يتجاوز %d" % f.get("max_absence", 0))
        if rep["include_tardy"]:
            conds.append("تأخر لا يتجاوز %d" % f.get("max_tardy", 0))
        if rep.get("excused_ok"):
            conds.append("الغياب بعذر لا يكسر الانضباط")
    cond_html = ("<div class='cond'>عوامل التصفية: %s</div>"
                 % _esc(" · ".join(conds))) if conds else ""

    cards = []
    if rep.get("kind") == "good":
        t = rep["totals"]
        cards.append(("عدد المنضبطين", t["rows"]))
        cards.append(("منهم بلا غياب ولا تأخر", t.get("perfect", 0)))
        if t.get("scanned"):
            cards.append(("من أصل", t["scanned"]))
            cards.append(("النسبة",
                          "%.0f%%" % (t["rows"] * 100.0 / t["scanned"])))
        cards.append(("أيام الدراسة في المدى", rep["school_days"]))
    else:
        if rep["include_absence"]:
            cards.append(("إجمالي أيام الغياب", rep["totals"]["absence"]))
        if rep["include_tardy"]:
            cards.append(("إجمالي أيام التأخر", rep["totals"]["tardy"]))
            cards.append(("مجموع دقائق التأخر",
                          _fmt_minutes(rep["totals"]["minutes"])))
        cards.append(("أيام الدراسة في المدى", rep["school_days"]))
        cards.append((_GROUP_AR.get(rep["group_by"], "بنود"),
                      rep["totals"]["rows"]))
    cards_html = "".join(
        "<div class='card'><div class='cv'>%s</div><div class='cl'>%s</div></div>"
        % (_esc(v), _esc(l)) for l, v in cards)

    printed = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    return """<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8">
<title>%(title)s</title><style>
  @page { size: A4 landscape; margin: 10mm; }
  *{box-sizing:border-box}
  body{font-family:"Segoe UI","Tahoma",Arial,sans-serif;margin:0;color:#111;
       background:#f5f6f7;font-size:12px}
  .page{background:#fff;padding:14mm;margin:8mm auto;max-width:277mm;
        box-shadow:0 0 6px rgba(0,0,0,.12)}
  .hd{display:flex;justify-content:space-between;align-items:flex-end;
      border-bottom:2px solid #15616d;padding-bottom:8px}
  .hd h1{margin:0;font-size:19px}
  .hd .sch{font-size:13px;color:#444;font-weight:600}
  .meta{margin:10px 0 4px;font-size:12px;color:#444}
  .cond{margin:4px 0 10px;font-size:12px;color:#8a4b12}
  .cards{display:flex;gap:8px;flex-wrap:wrap;margin:12px 0}
  .card{flex:1;min-width:110px;border:1px solid #d9e2db;border-radius:7px;
        padding:8px 10px;text-align:center;background:#fafbfa}
  .cv{font-size:18px;font-weight:700;color:#15616d}
  .cl{font-size:11px;color:#5d6b63;margin-top:2px}
  table{width:100%%;border-collapse:collapse;margin-top:6px}
  th,td{border:1px solid #cfd8d3;padding:5px 6px;text-align:center}
  th{background:#15616d;color:#fff;font-size:11px;font-weight:600}
  tbody tr:nth-child(even){background:#f7f9f8}
  td:nth-child(2){text-align:right}
  .empty{padding:22px;color:#777}
  .ft{margin-top:14px;padding-top:8px;border-top:1px solid #d9e2db;
      display:flex;justify-content:space-between;font-size:11px;color:#5d6b63}
  @media print{ body{background:#fff} .page{box-shadow:none;margin:0;
       padding:0;max-width:none} thead{display:table-header-group} }
</style></head><body><div class="page">
  <div class="hd"><h1>%(title)s</h1><div class="sch">%(school)s</div></div>
  <div class="meta">المدة: من <b>%(frm)s</b> إلى <b>%(to)s</b> ·
       التجميع حسب <b>%(grp)s</b></div>
  %(cond)s
  <div class="cards">%(cards)s</div>
  <table><thead><tr>%(head)s</tr></thead><tbody>%(body)s</tbody></table>
  <div class="ft"><span>طُبع في %(printed)s</span>
       <span>درب الطالب — darbstu.com</span></div>
</div></body></html>""" % {
        "title": _esc(title or "تقرير الغياب والتأخر"),
        "school": _esc(school_name or "المدرسة"),
        "frm": _esc(rep["from"]), "to": _esc(rep["to"]),
        "grp": _esc(_GROUP_AR.get(rep["group_by"], "")),
        "cond": cond_html, "cards": cards_html,
        "head": head, "body": "".join(body), "printed": printed,
    }


def build_rows_for_export(rep):
    """صفوفٌ مسطّحة للتصدير — نفس أعمدة الطباعة حرفاً، فلا يختلف الملف
    عن الورقة ويشكّ الوكيل في أيّهما الصحيح."""
    cols = columns(rep)
    out = [[h for h, _ in cols]]
    for i, r in enumerate(rep["rows"], 1):
        row = []
        for _, k in cols:
            if k == "idx":
                row.append(i)
            elif k in ("absence", "tardy", "minutes", "total", "excused"):
                row.append(int(r.get(k) or 0))      # أرقاماً لا نصّاً
            elif k in ("pct", "avg_minutes"):
                row.append(float(r.get(k) or 0))
            else:
                row.append(r.get(k, ""))
        out.append(row)
    return out


# ══════════════════════════════════════════════════════════════════
#  التقرير التفصيلي — سجلٌّ لكل فصل: الطالب صفّاً واليوم عموداً
# ══════════════════════════════════════════════════════════════════
# طلب ماهر (٢٦ أكتوبر ٢٠٢٦): «سجل كل فصل به كشف يعرض تأخر أو غياب كل طالب مع
# خانة للمجموع»، للتأخر وحده أو الغياب وحده أو لهما، يُطبع ويُحفظ.
# مصدر الصفوف **الكشف** لا السجلات: الطالب الذي لم يغب يظهر بصفٍّ فارغ،
# والطالب يُنسب لفصله في الكشف — سجلات البصمة لا تحمل معرّف فصل أصلاً.
DETAIL_MODES = ("both", "absence", "tardy")
DETAIL_TITLES = {"both": "التقرير التفصيلي — الغياب والتأخر",
                 "absence": "التقرير التفصيلي — الغياب",
                 "tardy": "التقرير التفصيلي — التأخر"}
_AR_DAYS = ("اثنين", "ثلاثاء", "أربعاء", "خميس", "جمعة", "سبت", "أحد")   # weekday(): الاثنين=0
_AR_MONTHS = ("يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو",
              "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر")


def build_detailed_register(start, end, students, mode="both", class_id=None, level=None):
    """
    `students`: الكشف مسطّحاً [{id, name, class_id, class_name}].
    يُرجع لكل فصل صفوف طلابه، وفي كل صفّ خانات الأيام ومجاميعه، ومجموع كل يوم.
    أيام الأعمدة: كل تاريخٍ سُجّل فيه غيابٌ أو تأخر **في المدرسة كلها** — فتتطابق
    أعمدة الفصول، والعطل لا سجلّ فيها فلا تظهر أعمدةً فارغة.
    """
    if mode not in DETAIL_MODES:
        mode = "both"
    want_abs, want_tdy = mode in ("both", "absence"), mode in ("both", "tardy")
    con = _con()
    try:
        cur = con.cursor()
        exempt = {str(r[0]) for r in cur.execute("SELECT student_id FROM exempted_students")}
        abs_rows = [(r["date"], str(r["student_id"])) for r in cur.execute(
            "SELECT date, student_id FROM absences WHERE date BETWEEN ? AND ?", [start, end])]
        tdy_rows = [(r["date"], str(r["student_id"]), int(r["mins"] or 0)) for r in cur.execute(
            "SELECT date, student_id, COALESCE(minutes_late,0) AS mins FROM tardiness "
            "WHERE date BETWEEN ? AND ?", [start, end])]
        excused = {(r["date"], str(r["student_id"])) for r in cur.execute(
            "SELECT date, student_id FROM excuses WHERE date BETWEEN ? AND ?", [start, end])}
    finally:
        con.close()

    days = sorted({d for d, _s in abs_rows} | {d for d, _s, _m in tdy_rows})
    ab, td = {}, {}
    for d, sid in abs_rows:
        ab.setdefault(sid, {})[d] = (d, sid) in excused
    for d, sid, m in tdy_rows:
        td.setdefault(sid, {})[d] = m

    classes, order, seen = {}, [], set()
    for s in students or []:
        sid = str(s.get("id") or "").strip()
        if not sid or sid in exempt or sid in seen:
            continue
        if not _cls_match(s.get("class_id"), class_id, level):
            continue
        seen.add(sid)
        key = _norm_cid(s.get("class_id")) or str(s.get("class_name") or "—")
        c = classes.get(key)
        if c is None:
            c = classes[key] = {"class_id": key, "class_name": s.get("class_name") or key,
                                "students": []}
            order.append(key)
        cells, n_abs, n_exc, n_tdy, mins = {}, 0, 0, 0, 0
        if want_abs:
            for d, exc in (ab.get(sid) or {}).items():
                cells.setdefault(d, {})["a"] = "x" if exc else "a"
                n_abs += 1
                n_exc += 1 if exc else 0
        if want_tdy:
            for d, m in (td.get(sid) or {}).items():
                cells.setdefault(d, {})["t"] = m
                n_tdy += 1
                mins += m
        c["students"].append({"id": sid, "name": s.get("name") or sid, "cells": cells,
                              "absence": n_abs, "excused": n_exc, "tardy": n_tdy,
                              "minutes": mins})

    out_classes = []
    for key in order:
        c = classes[key]
        c["students"].sort(key=lambda r: r["name"])
        c["day_absence"] = {d: sum(1 for r in c["students"] if "a" in r["cells"].get(d, {}))
                            for d in days}
        c["day_tardy"] = {d: sum(1 for r in c["students"] if "t" in r["cells"].get(d, {}))
                          for d in days}
        c["totals"] = {k: sum(r[k] for r in c["students"])
                       for k in ("absence", "excused", "tardy", "minutes")}
        c["totals"]["students"] = len(c["students"])
        out_classes.append(c)

    return {"ok": True, "kind": "detailed", "from": start, "to": end, "mode": mode,
            "title": DETAIL_TITLES[mode], "days": days, "school_days": len(days),
            "include_absence": want_abs, "include_tardy": want_tdy,
            "filters": {"class_id": class_id or "", "level": level or ""},
            "classes": out_classes,
            "totals": {k: sum(c["totals"][k] for c in out_classes)
                       for k in ("absence", "excused", "tardy", "minutes", "students")}}


def detailed_cell(cell):
    """نصّ الخانة: غ غياب · ع غياب بعذر · رقمٌ دقائق التأخر (ت بلا دقائق)."""
    if not cell:
        return ""
    parts = []
    if "a" in cell:
        parts.append("ع" if cell["a"] == "x" else "غ")
    if "t" in cell:
        parts.append(str(cell["t"]) if cell["t"] else "ت")
    return "/".join(parts)


def detailed_total_columns(rep):
    cols = []
    if rep["include_absence"]:
        cols += [("الغياب", "absence"), ("منها بعذر", "excused")]
    if rep["include_tardy"]:
        cols += [("التأخر", "tardy"), ("دقائق التأخر", "minutes")]
    return cols


def _day_head(d):
    dt = datetime.date.fromisoformat(d)
    return dt.day, _AR_DAYS[dt.weekday()]


def _month_groups(days):
    """[(اسم الشهر والسنة، عدد أيامه)] بترتيب الأيام — لصفّ العناوين الأعلى."""
    groups = []
    for d in days:
        dt = datetime.date.fromisoformat(d)
        label = "%s %d" % (_AR_MONTHS[dt.month - 1], dt.year)
        if groups and groups[-1][0] == label:
            groups[-1][1] += 1
        else:
            groups.append([label, 1])
    return groups


def build_detailed_print_html(rep, school_name=""):
    """كل فصلٍ في صفحة A4 أفقية. الخطّ يصغر كلما كثرت الأيام كي يتّسع الكشف."""
    days = rep["days"]
    tcols = detailed_total_columns(rep)
    n = len(days)
    fs = 10 if n <= 23 else (8.5 if n <= 35 else 7)
    legend = []
    if rep["include_absence"]:
        legend += ["غ = غياب", "ع = غياب بعذر"]
    if rep["include_tardy"]:
        legend += ["الرقم = دقائق التأخر", "ت = تأخر بلا دقائق"]

    pages = []
    for c in rep["classes"]:
        months = "".join("<th colspan='%d' class='mo'>%s</th>" % (k, _esc(lbl))
                         for lbl, k in _month_groups(days))
        dhead = "".join("<th class='dy'>%d<br><small>%s</small></th>" % _day_head(d) for d in days)
        thead = ("<tr><th rowspan='2' class='ix'>م</th><th rowspan='2' class='nm'>اسم الطالب</th>"
                 + months
                 + "".join("<th rowspan='2' class='tt'>%s</th>" % _esc(h) for h, _ in tcols)
                 + "</tr><tr>" + dhead + "</tr>")
        if not days:
            thead = ("<tr><th class='ix'>م</th><th class='nm'>اسم الطالب</th>"
                     + "".join("<th class='tt'>%s</th>" % _esc(h) for h, _ in tcols) + "</tr>")
        rows = []
        for i, r in enumerate(c["students"], 1):
            tds = "".join("<td class='%s'>%s</td>" % (
                ("a" if "a" in r["cells"].get(d, {}) else "t") if r["cells"].get(d) else "",
                _esc(detailed_cell(r["cells"].get(d)))) for d in days)
            tot = "".join("<td class='tv'>%s</td>" % _esc(r[k]) for _, k in tcols)
            rows.append("<tr><td>%d</td><td class='nm'>%s</td>%s%s</tr>"
                        % (i, _esc(r["name"]), tds, tot))
        if not rows:
            rows.append("<tr><td colspan='%d' class='empty'>لا طلاب في هذا الفصل</td></tr>"
                        % (2 + n + len(tcols)))

        def _day_sum(d):
            parts = []
            if rep["include_absence"]:
                parts.append(str(c["day_absence"][d]))
            if rep["include_tardy"]:
                parts.append(str(c["day_tardy"][d]))
            return "<br>".join(parts)
        foot = ("<tr class='sum'><td colspan='2'>المجموع%s</td>%s%s</tr>"
                % ("<br><small>غياب / تأخر</small>" if rep["mode"] == "both" else "",
                   "".join("<td>%s</td>" % _day_sum(d) for d in days),
                   "".join("<td class='tv'>%s</td>" % _esc(c["totals"][k]) for _, k in tcols)))
        pages.append("""<section class="page">
  <div class="hd"><h1>%(title)s</h1><div class="sch">%(school)s</div></div>
  <div class="meta">الفصل: <b>%(cls)s</b> · من <b>%(frm)s</b> إلى <b>%(to)s</b> ·
       أيام الدراسة: <b>%(nd)d</b> · عدد الطلاب: <b>%(ns)d</b></div>
  <div class="lg">%(legend)s</div>
  <table><thead>%(thead)s</thead><tbody>%(rows)s</tbody><tfoot>%(foot)s</tfoot></table>
  <div class="ft"><span>طُبع في %(printed)s</span><span>درب الطالب — darbstu.com</span></div>
</section>""" % {"title": _esc(rep["title"]), "school": _esc(school_name or "المدرسة"),
                 "cls": _esc(c["class_name"]), "frm": _esc(rep["from"]), "to": _esc(rep["to"]),
                 "nd": n, "ns": len(c["students"]), "legend": _esc(" · ".join(legend)),
                 "thead": thead, "rows": "".join(rows), "foot": foot,
                 "printed": datetime.datetime.now().strftime("%Y-%m-%d %H:%M")})
    if not pages:
        pages.append("<section class='page'><div class='empty'>لا فصول مطابقة لهذه الخيارات</div></section>")

    return """<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8">
<title>%(title)s</title><style>
  @page { size: A4 landscape; margin: 8mm; }
  *{box-sizing:border-box}
  body{font-family:"Segoe UI","Tahoma",Arial,sans-serif;margin:0;color:#111;background:#f5f6f7}
  .page{background:#fff;padding:10mm;margin:8mm auto;max-width:281mm;
        box-shadow:0 0 6px rgba(0,0,0,.12);page-break-after:always;break-after:page}
  .page:last-child{page-break-after:auto;break-after:auto}
  .hd{display:flex;justify-content:space-between;align-items:flex-end;
      border-bottom:2px solid #15616d;padding-bottom:6px}
  .hd h1{margin:0;font-size:17px}
  .hd .sch{font-size:13px;color:#444;font-weight:600}
  .meta{margin:8px 0 3px;font-size:12px;color:#444}
  .lg{margin:0 0 6px;font-size:11px;color:#5d6b63}
  table{width:100%%;border-collapse:collapse;font-size:%(fs)spx;table-layout:auto}
  th,td{border:1px solid #cfd8d3;padding:2px 1px;text-align:center;line-height:1.25}
  th{background:#15616d;color:#fff;font-weight:600}
  th.mo{background:#0f4a53}
  th.dy small{font-weight:400;opacity:.85;font-size:.82em}
  th.nm,td.nm{text-align:right;padding:2px 5px;white-space:nowrap}
  th.ix{width:22px}
  td.a{color:#b42318;font-weight:700}
  td.t{color:#1849a9;font-weight:600}
  td.tv{background:#f1f5f3;font-weight:700}
  th.tt{background:#25716f}
  tbody tr:nth-child(even){background:#f8faf9}
  tr.sum td{background:#e8f0ec;font-weight:700}
  .empty{padding:22px;color:#777;text-align:center}
  .ft{margin-top:8px;padding-top:6px;border-top:1px solid #d9e2db;
      display:flex;justify-content:space-between;font-size:10px;color:#5d6b63}
  @media print{ body{background:#fff} .page{box-shadow:none;margin:0;padding:0;max-width:none}
       thead{display:table-header-group} tr{page-break-inside:avoid} }
</style></head><body>%(pages)s</body></html>""" % {
        "title": _esc(rep["title"]), "fs": fs, "pages": "".join(pages)}


def _sheet_title(name, used):
    """اسم ورقة Excel: ٣١ حرفاً بلا []:*?/\\ — ولا يتكرر."""
    t = "".join("-" if ch in '[]:*?/\\' else ch for ch in str(name or "فصل")).strip()[:31] or "فصل"
    base, k = t, 2
    while t in used:
        suffix = " (%d)" % k
        t = base[:31 - len(suffix)] + suffix
        k += 1
    used.add(t)
    return t


def build_detailed_workbook(rep, school_name=""):
    """ملف Excel: ورقةٌ لكل فصل بأعمدة الطباعة نفسها — الخانات نصّاً والمجاميع أرقاماً."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    days, tcols = rep["days"], detailed_total_columns(rep)
    hdr = PatternFill("solid", fgColor="15616D")
    tot = PatternFill("solid", fgColor="E8F0EC")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    used = set()
    for c in rep["classes"] or [{"class_name": "لا فصول", "students": [], "day_absence": {},
                                 "day_tardy": {}, "totals": {}}]:
        ws = wb.create_sheet(_sheet_title(c["class_name"], used))
        ws.sheet_view.rightToLeft = True
        ws.append(["%s — %s" % (rep["title"], school_name or "المدرسة")])
        ws.append(["الفصل: %s · من %s إلى %s · أيام الدراسة: %d"
                   % (c["class_name"], rep["from"], rep["to"], len(days))])
        ws.append([])
        head = ["م", "اسم الطالب"] + ["%s\n%d %s" % (_AR_DAYS[datetime.date.fromisoformat(d).weekday()],
                                                    datetime.date.fromisoformat(d).day,
                                                    _AR_MONTHS[datetime.date.fromisoformat(d).month - 1])
                                      for d in days] + [h for h, _ in tcols]
        ws.append(head)
        for cell in ws[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = hdr
            cell.alignment = center
        for i, r in enumerate(c["students"], 1):
            ws.append([i, r["name"]] + [detailed_cell(r["cells"].get(d)) for d in days]
                      + [int(r[k]) for _, k in tcols])
        sums = []
        for d in days:
            parts = []
            if rep["include_absence"]:
                parts.append(str(c["day_absence"].get(d, 0)))
            if rep["include_tardy"]:
                parts.append(str(c["day_tardy"].get(d, 0)))
            sums.append(" / ".join(parts))
        ws.append(["", "المجموع" + (" (غياب / تأخر)" if rep["mode"] == "both" else "")] + sums
                  + [int(c["totals"].get(k, 0)) for _, k in tcols])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.fill = tot
        for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
            for cell in row:
                if cell.column != 2:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions["A"].width = 5
        ws.column_dimensions["B"].width = max(14, min(36, max((len(r["name"]) for r in c["students"]),
                                                             default=10) + 2))
        for j in range(len(days)):
            ws.column_dimensions[get_column_letter(3 + j)].width = 6.5
        for j in range(len(tcols)):
            ws.column_dimensions[get_column_letter(3 + len(days) + j)].width = 11
        ws.row_dimensions[4].height = 34
        ws.freeze_panes = "C5"
    import io as _io
    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
